"""Flask integration tests: every page and workflow through real HTTP requests."""
import io
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from config import Config
from opportunity_finder.app_factory import create_app

GOOD = dict(asin="B0ROUTE001", title="Route test kettle", brand="Kettleco", fee_category="everything_else",
            selling_price="24.99", bsr="5000", monthly_sales="300", total_sellers="6", fba_sellers="3",
            sourcing_price="8.00", shipping_prep="0.50", other_costs="0", weight_g="400", length_cm="20",
            width_cm="15", height_cm="8", fulfilment="FBA", restriction_status="ungated")


def make_app(data_dir: Path, **overrides):
    options = {"SECRET_KEY": "test-secret", "TESTING": True, "CSRF_ENABLED": False, **overrides}
    cfg = Config(DATA_DIR=data_dir, **options)
    return create_app(cfg)


@pytest.fixture
def data_dir(tmp_path):
    return tmp_path / "data"


@pytest.fixture
def client(data_dir):
    return make_app(data_dir).test_client()


def add(client, **overrides):
    return client.post("/products/new", data={**GOOD, **overrides})


def settle(client) -> None:
    """Render one page so pending flash toasts (which name products) don't leak into the next assertion."""
    client.get("/")


def product_count(client) -> int:
    html = client.get("/products?view=all").get_data(as_text=True)
    return len(re.findall(r'name="ids" value="\d+"', html))


def text(response) -> str:
    return response.get_data(as_text=True)


# ------------------------------------------------------------------ pages
@pytest.mark.parametrize("url", ["/", "/products", "/products/new", "/products/new?method=manual", "/opportunities",
                                 "/rejected", "/rejected?tab=incomplete", "/import", "/settings", "/health"])
def test_every_page_renders_when_empty(client, url):
    response = client.get(url)
    assert response.status_code == 200
    if url == "/":
        assert "No products yet" in text(response)


def test_security_headers_and_no_external_assets(client):
    response = client.get("/")
    assert "default-src 'self'" in response.headers["Content-Security-Policy"]
    assert response.headers["X-Frame-Options"] == "DENY"
    html = text(response)
    assert not re.search(r'<(script|link)[^>]+(src|href)="https?://', html)


# ------------------------------------------------------- add / edit / delete
def test_add_edit_delete_flow(client):
    response = add(client)
    assert response.status_code == 302 and response.headers["Location"].endswith("/products/1")
    detail = text(client.get("/products/1"))
    assert "Qualified because" in detail and "£9.51" in detail and "Profit breakdown" in detail
    assert "Confirm exact fees in Amazon Seller Central Revenue Calculator before purchasing stock." in detail

    assert client.get("/products/1/edit").status_code == 200
    response = client.post("/products/1/edit", data={**GOOD, "bsr": "42831", "monthly_sales": "34"})
    assert response.status_code == 302
    detail = text(client.get("/products/1"))
    assert "Rejected because" in detail
    assert "BSR 42,831 exceeds maximum 25,000" in detail and "Monthly sales 34 below minimum 50" in detail

    confirm_page = client.post("/products/1/delete", data={})
    assert confirm_page.status_code == 200 and "Delete this product?" in text(confirm_page)
    assert client.post("/products/1/delete", data={"confirm": "yes"}).status_code == 302
    missing = client.get("/products/1")
    assert missing.status_code == 404 and "Page not found" in text(missing) and "Traceback" not in text(missing)


def test_backend_validation_blocks_bad_data(client):
    response = add(client, selling_price="-5", total_sellers="2.5", asin="NOPE", amazon_url="not a url")
    body = text(response)
    assert response.status_code == 422
    assert "cannot be negative" in body and "whole number" in body and "10 letters and numbers" in body
    assert product_count(client) == 0


def test_missing_data_saves_as_incomplete(client):
    assert add(client, bsr="", monthly_sales="").status_code == 302
    detail = text(client.get("/products/1"))
    assert "Incomplete data" in detail and "Best Sellers Rank is missing" in detail


def test_duplicate_detection_choices(client):
    add(client)
    clash = add(client, selling_price="26.00")
    assert clash.status_code == 409 and "This looks like a product you already have" in text(clash)
    assert product_count(client) == 1

    cancel = client.post("/products/new", data={**GOOD, "duplicate_action": "cancel"})
    assert cancel.status_code == 302 and product_count(client) == 1

    update = client.post("/products/new", data={**GOOD, "selling_price": "26.00", "duplicate_action": "update", "duplicate_id": "1"})
    assert update.headers["Location"].endswith("/products/1")
    assert "£26.00" in text(client.get("/products/1")) and product_count(client) == 1

    create = client.post("/products/new", data={**GOOD, "duplicate_action": "create"})
    assert create.status_code == 302 and product_count(client) == 2


def test_smart_paste_prefills_without_saving(client):
    fixture = (Path(__file__).parent / "fixtures" / "amazon_product_page.txt").read_text(encoding="utf-8")
    response = client.post("/products/new/paste", data={"pasted_text": fixture})
    body = text(response)
    assert response.status_code == 200
    assert "Auto extracted — please verify." in body
    assert 'value="B0TESTAB12"' in body and 'value="23.99"' in body and 'value="3456"' in body
    assert "Several prices found" in body
    assert product_count(client) == 0

    nothing = client.post("/products/new/paste", data={"pasted_text": "hello there"})
    assert nothing.status_code == 422 and "Nothing recognisable was found" in text(nothing)


def test_live_preview_api(client):
    ok = client.post("/api/preview", data=GOOD).get_json()
    assert ok["status"] == "qualified" and ok["profit"]["net"] == "£9.51" and ok["score"] == 91
    missing = client.post("/api/preview", data={**GOOD, "selling_price": ""}).get_json()
    assert missing["status"] == "incomplete"
    assert "Amazon selling price is required before profitability can be calculated." in missing["missing"]
    bad = client.post("/api/preview", data={**GOOD, "sourcing_price": "abc"})
    assert bad.status_code == 200 and "sourcing_price" in bad.get_json()["errors"]


# ------------------------------------------------------------ list & bulk
def test_search_filter_sort_paginate(client):
    add(client)
    add(client, asin="B0ROUTE002", title="Rejected teapot", bsr="90000")
    add(client, asin="B0ROUTE003", title="Risky mug", flag_ip_concern="1", selling_price="30.00")
    settle(client)

    rejected = text(client.get("/products?status=rejected"))
    assert "Rejected teapot" in rejected and "Route test kettle" not in rejected
    assert "Risky mug" in text(client.get("/products?risk=HIGH")) and "Route test kettle" not in text(client.get("/products?risk=HIGH"))
    searched = text(client.get("/products?q=route+test"))  # not "kettle": every test product's brand is "Kettleco"
    assert "Route test kettle" in searched and "Risky mug" not in searched
    by_price = text(client.get("/products?sort=price&dir=desc"))
    assert by_price.index("Risky mug") < by_price.index("Route test kettle")
    assert "Route test kettle" not in text(client.get("/products?bsr_min=6000"))
    assert "must be a number" in text(client.get("/products?roi_min=abc"))
    assert client.get("/products?per_page=10&page=99").status_code == 200


def test_bulk_archive_and_delete(client):
    add(client)
    add(client, asin="B0ROUTE002", title="Second")
    client.post("/products/bulk", data={"ids": ["2"], "action": "archive"})
    settle(client)
    assert "Second" not in text(client.get("/products")) and "Second" in text(client.get("/products?view=archived"))

    client.post("/products/bulk", data={"ids": ["1", "2"], "action": "delete"})
    assert product_count(client) == 2  # no confirmation -> nothing deleted
    client.post("/products/bulk", data={"ids": ["1", "2"], "action": "delete", "confirm": "yes"})
    assert product_count(client) == 0


def test_duplicate_archive_reevaluate_routes(client):
    add(client)
    response = client.post("/products/1/duplicate")
    assert response.headers["Location"].endswith("/products/2/edit")
    assert "(copy)" in text(client.get("/products/2"))
    client.post("/products/2/archive")
    assert "This product is archived" in text(client.get("/products/2"))
    client.post("/products/2/unarchive")
    assert "This product is archived" not in text(client.get("/products/2"))
    assert client.post("/products/1/reevaluate", follow_redirects=True).status_code == 200
    assert client.post("/products/reevaluate-all", follow_redirects=True).status_code == 200


def test_opportunities_and_rejected_pages(client):
    add(client)
    add(client, asin="B0ROUTE002", title="Rejected teapot", bsr="90000")
    add(client, asin="B0ROUTE003", title="No sales data", monthly_sales="")
    opp = text(client.get("/opportunities"))
    assert "Route test kettle" in opp and "Target profit plan" in opp and "Short by" in opp
    assert "Rejected teapot" in text(client.get("/rejected"))
    assert "No sales data" in text(client.get("/rejected?tab=incomplete"))
    dash = text(client.get("/"))
    assert "Top opportunities" in dash and "Route test kettle" in dash


# ---------------------------------------------------------------- settings
def settings_form(**changes):
    from opportunity_finder.domain.settings import Settings
    data = {k: ("1" if v is True else str(v)) for k, v in Settings().to_dict().items() if v is not False}
    data.update(changes)
    return data


def test_settings_save_reevaluate_reset_and_persist(client, data_dir):
    add(client, bsr="42831")
    assert "Rejected because" in text(client.get("/products/1"))

    assert client.post("/settings", data=settings_form(max_bsr="50000")).status_code == 302
    assert 'value="50000"' in text(client.get("/settings"))
    assert "Qualified because" in text(client.get("/products/1"))

    restarted = make_app(data_dir).test_client()
    assert 'value="50000"' in text(restarted.get("/settings"))
    assert "Qualified because" in text(restarted.get("/products/1"))

    bad = client.post("/settings", data=settings_form(max_bsr="-1", capture_rate_low="150"))
    assert bad.status_code == 422 and "must be at least" in text(bad)

    client.post("/settings/reset")
    assert 'value="25000"' in text(client.get("/settings"))


# ------------------------------------------------------------------ import
CSV = ("ASIN,Title,Amazon Price,Sales Rank,Monthly Sales,Sellers,FBA Sellers,Cost,Weight (g)\n"
       "B0IMPORT01,Imported lamp,24.99,5000,300,6,3,8.00,400\n"
       "B0IMPORT02,Broken row,-3,5000,300,6,3,8.00,400\n")


def test_import_preview_and_confirm(client):
    upload = client.post("/import", data={"file": (io.BytesIO(CSV.encode()), "products.csv")}, content_type="multipart/form-data")
    assert upload.status_code == 302
    preview_url = upload.headers["Location"]
    preview = text(client.get(preview_url))
    assert "Check your import" in preview and "cannot be negative" in preview

    token_mapping = {f"map_{i}": key for i, key in enumerate(
        ["asin", "title", "selling_price", "bsr", "monthly_sales", "total_sellers", "fba_sellers", "sourcing_price", "weight_g"])}
    done = client.post(preview_url, data={**token_mapping, "policy": "skip", "step": "confirm"})
    body = text(done)
    assert done.status_code == 200 and "Import finished" in body and "Rows not imported" in body
    assert product_count(client) == 1
    assert client.get(preview_url, follow_redirects=True).status_code == 200  # session gone -> back to start


def test_import_errors_are_friendly(client):
    empty = client.post("/import", data={"file": (io.BytesIO(b""), "empty.csv")}, content_type="multipart/form-data")
    assert empty.status_code == 400 and "The file is empty." in text(empty)
    nothing = client.post("/import", data={})
    assert nothing.status_code == 400 and "Choose a CSV or Excel file" in text(nothing)
    pasted = client.post("/import", data={"pasted": "ASIN\tTitle\nB0PASTE001\tDesk lamp\n"})
    assert pasted.status_code == 302
    assert "expired" in text(client.get("/import/" + "0" * 32, follow_redirects=True))


def test_template_downloads(client):
    csv_file = client.get("/import/template.csv")
    assert csv_file.status_code == 200 and "ASIN,Title" in csv_file.get_data().decode("utf-8-sig")
    xlsx = client.get("/import/template.xlsx")
    assert load_workbook(io.BytesIO(xlsx.data)).sheetnames == ["Products", "Instructions", "Fee categories"]


# ------------------------------------------------------------------ export
def test_exports(client):
    add(client)
    add(client, asin="B0ROUTE002", title="Rejected teapot", bsr="90000")
    response = client.get("/export/excel")
    assert response.status_code == 200 and "spreadsheetml" in response.content_type
    assert "attachment" in response.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames[:4] == ["Qualified", "Rejected", "Incomplete", "Summary"]
    filtered = load_workbook(io.BytesIO(client.get("/export/excel?scope=filtered&status=rejected").data))
    assert filtered["Qualified"]["A2"].value == "No qualified products yet."
    selected = client.post("/products/bulk", data={"ids": ["1"], "action": "export_csv"})
    assert "/export/csv" in selected.headers["Location"]
    csv_text = client.get(selected.headers["Location"]).get_data().decode("utf-8-sig")
    assert "Route test kettle" in csv_text and "Rejected teapot" not in csv_text


# ------------------------------------------------------------------ backups
def test_backup_and_restore_routes(client):
    add(client)
    backup = client.post("/settings/backup")
    assert backup.status_code == 200 and backup.data.startswith(b"SQLite format 3")
    client.post("/products/1/delete", data={"confirm": "yes"})
    assert product_count(client) == 0

    no_confirm = client.post("/settings/restore", data={"backup": (io.BytesIO(backup.data), "b.db")}, content_type="multipart/form-data")
    assert no_confirm.status_code == 302 and product_count(client) == 0
    client.post("/settings/restore", data={"backup": (io.BytesIO(backup.data), "b.db"), "confirm": "yes"}, content_type="multipart/form-data")
    assert product_count(client) == 1
    junk = client.post("/settings/restore", data={"backup": (io.BytesIO(b"junk"), "b.db"), "confirm": "yes"},
                       content_type="multipart/form-data", follow_redirects=True)
    assert "not an Opportunity Finder backup" in text(junk)


# ----------------------------------------------------------------- safety
def test_csrf_and_host_checks(data_dir):
    client = make_app(data_dir, CSRF_ENABLED=True).test_client()
    blocked = client.post("/products/new", data=GOOD)
    assert blocked.status_code == 400 and "Reload the page" in text(blocked)
    token = re.search(r'name="csrf-token" content="([^"]+)"', text(client.get("/products/new"))).group(1)
    assert client.post("/products/new", data={**GOOD, "csrf_token": token}).status_code == 302
    assert client.post("/api/preview", data=GOOD, headers={"X-CSRF-Token": token}).status_code == 200
    assert client.get("/", headers={"Host": "evil.example.com"}).status_code == 400


def test_unexpected_errors_are_friendly(client, monkeypatch):
    from opportunity_finder.routes import dashboard

    def boom():
        raise RuntimeError("database exploded")

    monkeypatch.setattr(dashboard, "get_service", boom)
    response = client.get("/")
    body = text(response)
    assert response.status_code == 500 and "Something went wrong" in body and "Reference" in body
    assert "database exploded" not in body and "Traceback" not in body
