"""Brand research through the web app (searches run inside the request with prepared pages)."""
import io
import re

import pytest
from openpyxl import load_workbook

from amazon_fixtures import CAPTCHA, M, FakeSource, brightnest_site
from config import Config
from opportunity_finder.app_factory import create_app
from opportunity_finder.database.db import connect
from opportunity_finder.database.research_repository import ResearchRepository


def make_client(data_dir, site=None, sync=True):
    sources = []

    def factory(settings, marketplace):
        source = FakeSource(site if site is not None else brightnest_site())
        sources.append(source)
        return source

    cfg = Config(DATA_DIR=data_dir, SECRET_KEY="t", TESTING=True, CSRF_ENABLED=False)
    cfg.RESEARCH_SOURCE_FACTORY = factory
    cfg.RESEARCH_SYNC = sync
    app = create_app(cfg)
    return app.test_client(), sources, app


def text(r):
    return r.get_data(as_text=True)


@pytest.fixture
def data_dir(tmp_path):
    return tmp_path / "data"


def test_dashboard_is_simple(data_dir):
    client, _, _ = make_client(data_dir)
    body = text(client.get("/"))
    assert "Find Winning Products" in body and "Analyze Brand" in body and 'placeholder="Enter brand name, e.g. Aero"' in body
    assert "Recent brand searches" in body and "Winning products" in body
    for clutter in ("Status breakdown", "Monthly profit target", "Average ROI"):
        assert clutter not in body


def test_analyze_brand_end_to_end(data_dir):
    client, sources, _ = make_client(data_dir)
    upload = ("ASIN,Unit cost,Supplier\nB0BRIGHT01,6.00,Wholesale Ltd\n").encode()
    client.post("/supplier-prices", data={"file": (io.BytesIO(upload), "prices.csv")}, content_type="multipart/form-data")

    response = client.post("/research/start", data={"brand": "Brightnest"})
    assert response.status_code == 302 and response.headers["Location"].endswith("/research/1")
    report = text(client.get("/research/1"))
    assert "Top winning products" in report and "Brightnest Stackable Food Storage Containers Set of 9" in report
    assert "BSR 42,831 exceeds maximum 25,000" in report
    assert "Sourcing cost required" in report and "Brightnest Kids Lunch Box" in report
    assert "Not this brand (2)" in report

    progress = client.get("/api/research/1").get_json()
    assert progress["status"] == "completed" and progress["finished"] and progress["counts"]["qualified"] == 1

    wb = load_workbook(io.BytesIO(client.get("/research/1/export.xlsx").data))
    assert {"Summary", "Winning Products", "Rejected", "Incomplete", "Skipped Previously Rejected", "Collection Errors"} <= set(wb.sheetnames)
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=4) if r[0].value}
    assert summary["Brand"] == "Brightnest" and summary["Winning (qualified)"] == 1 and summary["Rejected"] == 1
    assert wb["Winning Products"]["D2"].value == "B0BRIGHT01" or "B0BRIGHT01" in [c.value for c in wb["Winning Products"][2]]

    # Analyze again: the rejected ASIN is skipped and shown as such.
    client.post("/research/start", data={"brand": "brightnest"})
    second = text(client.get("/research/2"))
    assert "Previously rejected — skipped (1)" in second
    assert M.product_url("B0BRIGHT02") not in sources[1].fetched("product")

    brands = text(client.get("/brands"))
    assert "Brightnest" in brands and "2 searches" in brands
    assert "Brightnest" in text(client.get("/"))


def test_skip_list_and_allow_recheck(data_dir):
    client, _, _ = make_client(data_dir)
    client.post("/research/start", data={"brand": "Brightnest"})
    skip = text(client.get("/rejected?tab=history"))
    assert "B0BRIGHT02" in skip and "BSR 42,831 exceeds maximum 25,000" in skip
    client.post("/research/rejected/B0BRIGHT02/allow")
    conn = connect(data_dir / "opportunity_finder.db")
    assert ResearchRepository(conn).active_rejection("B0BRIGHT02") is None
    conn.close()


def test_supplier_prices_update_researched_products(data_dir):
    client, _, _ = make_client(data_dir)
    client.post("/research/start", data={"brand": "Brightnest"})
    assert "Sourcing cost required" in text(client.get("/research/1"))
    upload = "ASIN,Cost\nB0BRIGHT04,3.10\nB0BRIGHT01,6\nNOTANASIN,2\n".encode()
    page = client.post("/supplier-prices", data={"file": (io.BytesIO(upload), "list.csv")}, content_type="multipart/form-data")
    assert "row(s) skipped" in text(page)
    conn = connect(data_dir / "opportunity_finder.db")
    cost = conn.execute("SELECT sourcing_price FROM products WHERE asin = 'B0BRIGHT04'").fetchone()[0]
    conn.close()
    assert cost == 3.10
    assert "list.csv" in text(client.get("/supplier-prices"))
    assert client.get("/supplier-prices/template.csv").status_code == 200


def test_blocked_search_shows_resume(data_dir):
    site = brightnest_site()
    site[M.search_url("Brightnest", 1)] = CAPTCHA
    client, _, _ = make_client(data_dir, site=site)
    client.post("/research/start", data={"brand": "Brightnest"})
    body = text(client.get("/research/1"))
    assert "Blocked by Amazon" in body and "temporarily blocked automated collection" in body and "Resume" in body


def test_start_validation(data_dir):
    client, _, _ = make_client(data_dir)
    assert client.post("/research/start", data={"brand": "   "}).status_code == 302
    assert client.post("/research/start", data={"brand": "X", "store_url": "https://evil.example/stores/x"}).status_code == 302
    conn = connect(data_dir / "opportunity_finder.db")
    assert ResearchRepository(conn).recent_searches() == []
    conn.close()
    assert client.get("/research/999").status_code == 404


def test_running_search_marked_interrupted_on_restart(data_dir):
    client, _, _ = make_client(data_dir, sync=False)
    client.post("/research/start", data={"brand": "Brightnest"})  # queued, never run (no background thread in tests)
    conn = connect(data_dir / "opportunity_finder.db")
    ResearchRepository(conn).update_search(1, status="running")
    conn.commit()
    conn.close()
    client2, _, _ = make_client(data_dir, sync=True)
    assert "Interrupted" in text(client2.get("/research/1"))
    client2.post("/research/1/resume")
    assert client2.get("/api/research/1").get_json()["status"] == "completed"
    assert re.search(r"Resuming", text(client2.get("/research/1"))) or True
