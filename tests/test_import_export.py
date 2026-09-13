"""Bulk import (CSV / XLSX / pasted TSV) and Excel / CSV export."""
import io

import pytest
from openpyxl import load_workbook

from conftest import TODAY, good_product
from opportunity_finder.constants import QUALIFIED
from opportunity_finder.database.db import connect, init_db
from opportunity_finder.domain.settings import Settings
from opportunity_finder.services import import_service as imp
from opportunity_finder.services.excel_export import build_csv, build_workbook, safe_text
from opportunity_finder.services.product_service import ProductService
from opportunity_finder.services.target_plan import build_target_plan

HEADER = "ASIN,Title,Brand,Amazon Price,Sales Rank,Monthly Sales,Sellers,FBA Sellers,Cost,Weight (g),Length (cm),Width (cm),Height (cm),Restriction,Amazon sells\n"


@pytest.fixture
def service(tmp_path):
    path = tmp_path / "t.db"
    init_db(path)
    conn = connect(path)
    yield ProductService(conn, today=lambda: TODAY)
    conn.close()


def csv_bytes(*rows: str) -> bytes:
    return (HEADER + "".join(r + "\n" for r in rows)).encode("utf-8")


GOOD_ROW = 'B0IMPORT01,Garden Kneeler,GreenCo,"£1,024.99",5000,300,6,3,£8.00,400,20,15,8,ungated,no'


# ------------------------------------------------------------------ reading
def test_csv_read_and_auto_map():
    table = imp.read_csv_bytes(csv_bytes(GOOD_ROW))
    mapping = imp.auto_map(table.headers)
    assert set(mapping.values()) >= {"asin", "title", "brand", "selling_price", "bsr", "monthly_sales", "total_sellers",
                                     "fba_sellers", "sourcing_price", "weight_g", "restriction_status", "flag_amazon_sells"}
    assert table.rows[0][3] == "£1,024.99"


@pytest.mark.parametrize("data,message", [
    (b"", "empty"),
    (b"   \n\n", "empty"),
    (HEADER.encode(), "No data rows"),
    (b"\x00\x01\x02binary", "doesn't look like a CSV"),
])
def test_bad_csv_files(data, message):
    with pytest.raises(imp.ImportFileError, match=message):
        imp.read_csv_bytes(data)


def test_upload_type_checks():
    with pytest.raises(imp.ImportFileError, match=".xls files"):
        imp.read_upload("old.xls", b"x")
    with pytest.raises(imp.ImportFileError, match="Upload a .csv or .xlsx"):
        imp.read_upload("notes.pdf", b"%PDF")
    with pytest.raises(imp.ImportFileError, match="valid .xlsx"):
        imp.read_upload("broken.xlsx", b"PK not really")


def test_pasted_tab_separated():
    text = "ASIN\tTitle\tAmazon price\nB0PASTE001\tDesk Lamp\t19.99\n\nB0PASTE002\tFloor Lamp\t49.00\n"
    table = imp.read_delimited_text(text)
    assert table.headers == ["ASIN", "Title", "Amazon price"]
    assert len(table.rows) == 2 and table.blank_rows == 1


def test_templates_round_trip():
    csv_table = imp.read_csv_bytes(imp.template_csv(), "template.csv")
    assert len(imp.auto_map(csv_table.headers)) == len(imp.TEMPLATE_FIELDS)
    xlsx_table = imp.read_upload("template.xlsx", imp.template_xlsx())
    assert xlsx_table.headers == csv_table.headers
    assert xlsx_table.rows[0][0] == "B0EXAMPLE1"


def test_too_many_rows():
    data = HEADER + ("B0IMPORT01,x,,1,1,1,1,1,1,1,1,1,1,,\n" * (imp.MAX_ROWS + 1))
    with pytest.raises(imp.ImportFileError, match="at most"):
        imp.read_csv_bytes(data.encode())


# ------------------------------------------------------------------ preview
def test_preview_validates_every_row(service):
    rows = [
        GOOD_ROW,
        "B0IMPORT02,Negative price,,-5,5000,300,6,3,8,400,20,15,8,,",
        "B0IMPORT03,Half seller,,20,5000,300,6.5,3,8,400,20,15,8,,",
        "BADASIN,Bad asin,,20,5000,300,6,3,8,400,20,15,8,,",
        "B0IMPORT04,Too many FBA,,20,5000,300,2,3,8,400,20,15,8,,",
        "B0IMPORT05,Weird restriction,,20,5000,300,6,3,8,400,20,15,8,maybe,",
        "B0IMPORT01,Same as first row,,20,5000,300,6,3,8,400,20,15,8,,",
        ",Missing lots,,,,,,,,,,,,,",
    ]
    table = imp.read_csv_bytes(csv_bytes(*rows))
    preview = imp.build_preview(table, imp.auto_map(table.headers), service.repo)
    by_row = {r.row_number: r for r in preview.rows}
    assert by_row[2].action == "import" and not by_row[2].errors
    assert "cannot be negative" in by_row[3].errors[0]
    assert "whole number" in by_row[4].errors[0]
    assert "10 letters and numbers" in by_row[5].errors[0]
    assert "can't be more than total sellers" in by_row[6].errors[0]
    assert "not recognised" in by_row[7].errors[0]
    assert "Same product as row 2" in by_row[8].errors[0]
    assert by_row[9].action == "import"  # sparse but identifiable -> imported as Incomplete Data
    assert preview.counts == {"import": 2, "update": 0, "skip": 0, "invalid": 6, "total": 8}


def test_mapping_requires_identity(service):
    table = imp.read_delimited_text("Amazon price\tBSR\n9.99\t100\n")
    preview = imp.build_preview(table, imp.auto_map(table.headers), service.repo)
    assert preview.mapping_errors and all(r.action == "invalid" for r in preview.rows)
    with pytest.raises(imp.ImportFileError):
        imp.commit_preview(preview, service)


def test_mapping_from_form_and_double_mapping():
    headers = ["A", "B"]
    mapping = imp.mapping_from_form({"map_0": "asin", "map_1": "asin"}, headers)
    assert any("more than one column" in e for e in imp.check_mapping(mapping))
    assert imp.mapping_from_form({"map_0": "not_a_field"}, headers) == {}


@pytest.mark.parametrize("policy,expected_action", [("skip", "skip"), ("update", "update"), ("create", "import")])
def test_duplicate_policies(service, policy, expected_action):
    existing = good_product(asin="B0IMPORT01")
    existing.id = None
    service.create(existing)
    table = imp.read_csv_bytes(csv_bytes(GOOD_ROW))
    preview = imp.build_preview(table, imp.auto_map(table.headers), service.repo, policy)
    assert preview.rows[0].action == expected_action
    assert preview.rows[0].duplicate.matched_on == "asin"


def test_commit_creates_updates_and_reports(service):
    existing = good_product(asin="B0IMPORT01", sourcing_price=5.0, supplier="Keep Supplier")
    existing.id = None
    old_id = service.create(existing)
    rows = [GOOD_ROW, "B0IMPORT09,New thing,,20,5000,300,6,3,8,400,20,15,8,,yes", "BAD,Bad,,20,,,,,,,,,,,"]
    table = imp.read_csv_bytes(csv_bytes(*rows))
    preview = imp.build_preview(table, imp.auto_map(table.headers), service.repo, "update")
    report = imp.commit_preview(preview, service)
    assert (report.imported, report.updated, report.skipped, len(report.failed)) == (1, 1, 0, 1)
    assert report.failed[0][0] == 4

    updated = service.get(old_id)
    assert updated.selling_price == 1024.99 and updated.supplier == "Keep Supplier" and updated.source == "import"
    created = service.get(report.new_ids[0])
    assert created.flag_amazon_sells is True and created.source == "import"
    assert created.evaluation["status"] in {QUALIFIED, "rejected", "high_risk", "needs_verification", "incomplete"}


def test_commit_rolls_back_on_database_error(service, monkeypatch):
    rows = [GOOD_ROW, "B0IMPORT09,Second,,20,5000,300,6,3,8,400,20,15,8,,"]
    table = imp.read_csv_bytes(csv_bytes(*rows))
    preview = imp.build_preview(table, imp.auto_map(table.headers), service.repo)
    real_create = service.create
    calls = {"n": 0}

    def flaky(product, source=None, commit=True):
        calls["n"] += 1
        if calls["n"] == 2:
            raise RuntimeError("disk full")
        return real_create(product, source=source, commit=commit)

    monkeypatch.setattr(service, "create", flaky)
    with pytest.raises(RuntimeError):
        imp.commit_preview(preview, service)
    assert service.repo.count(include_archived=True) == 0


def test_import_weight_in_kg(service):
    table = imp.read_delimited_text("ASIN\tWeight (kg)\nB0KILOS001\t1.25\n")
    preview = imp.build_preview(table, imp.auto_map(table.headers), service.repo)
    assert preview.rows[0].product.weight_g == 1250


def test_session_storage(tmp_path):
    table = imp.read_delimited_text("ASIN\nB0SESSION1\n")
    token = imp.save_session(tmp_path, table)
    assert imp.load_session(tmp_path, token).rows == [["B0SESSION1"]]
    assert imp.load_session(tmp_path, "../../etc/passwd") is None
    imp.delete_session(tmp_path, token)
    assert imp.load_session(tmp_path, token) is None


# ------------------------------------------------------------------ export
def _items(service):
    for kw in [dict(asin="B0QUAL0001", title="Winner"), dict(asin="B0REJ00001", title="Loser", bsr=90_000),
               dict(asin="B0INC00001", title="Unknown sales", monthly_sales=None),
               dict(asin="B0RISK0001", title="=HYPERLINK(\"http://evil\")", flag_ip_concern=True)]:
        p = good_product(**kw)
        p.id = None
        service.create(p)
    return service.items()


def test_excel_workbook(service):
    items = _items(service)
    plan = build_target_plan(items, Settings())
    data = build_workbook(items, plan, Settings())
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[:4] == ["Qualified", "Rejected", "Incomplete", "Summary"]
    assert {"Needs Review", "Target Plan", "Assumptions"} <= set(wb.sheetnames)

    q = wb["Qualified"]
    headers = [c.value for c in q[1]]
    for required in ["ASIN", "Product", "Brand", "Category", "Amazon price", "BSR", "Sellers", "FBA sellers",
                     "Estimated monthly sales (EST)", "Sourcing price", "Supplier", "Referral fee (EST)", "FBA fee (EST)",
                     "Shipping/prep", "Other costs", "Net profit/unit (EST)", "ROI (EST)", "Conservative units/month (EST)",
                     "Conservative monthly profit (EST)", "Theoretical monthly profit (EST)", "Restriction status", "Risk",
                     "Risk reasons", "Included in target plan", "Source URL", "Checked date"]:
        assert required in headers, required
    assert q.freeze_panes == "E2" and q.auto_filter.ref.startswith("A1:")
    assert q["A1"].font.bold
    row = {h: q.cell(2, i + 1) for i, h in enumerate(headers)}
    assert row["Product"].value == "Winner"
    assert row["Amazon price"].number_format.startswith("£")
    assert row["ROI (EST)"].number_format == "0.0%" and row["ROI (EST)"].value == pytest.approx(1.1888, abs=1e-3)
    assert row["Included in target plan"].value == "Yes"
    assert q.max_row == 2

    rejected_headers = [c.value for c in wb["Rejected"][1]]
    reasons = wb["Rejected"].cell(2, rejected_headers.index("Rejected because") + 1).value
    assert "BSR 90,000 exceeds maximum 25,000" in reasons
    incomplete_headers = [c.value for c in wb["Incomplete"][1]]
    assert "monthly sales" in wb["Incomplete"].cell(2, incomplete_headers.index("Missing data") + 1).value.lower()

    review = wb["Needs Review"]
    title_cell = review.cell(2, [c.value for c in review[1]].index("Product") + 1)
    assert title_cell.data_type == "s" and title_cell.value.startswith("'=")

    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=4) if r[0].value}
    assert summary["Product count"] == 4 and summary["Qualified count"] == 1
    assert summary["Rejected count"] == 1 and summary["Incomplete count"] == 1
    assert summary["Target achieved?"] == "NO"
    assert summary["Monthly profit target"] == 1000


def test_excel_with_no_products():
    data = build_workbook([], build_target_plan([], Settings()), Settings())
    wb = load_workbook(io.BytesIO(data))
    assert wb["Qualified"]["A2"].value == "No qualified products yet."


def test_csv_export(service):
    items = _items(service)
    data = build_csv(items)
    assert data.startswith("﻿".encode())
    text = data.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Status,Score,ASIN")
    assert "'=HYPERLINK" in text


def test_safe_text():
    assert safe_text("=1+1") == "'=1+1" and safe_text("@SUM") == "'@SUM" and safe_text("Normal") == "Normal"
    assert safe_text(12) == 12
