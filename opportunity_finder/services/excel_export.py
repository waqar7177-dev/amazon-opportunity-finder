"""Excel workbook and CSV export.

Sheets: Qualified, Rejected, Incomplete, Summary (required) plus Needs Review,
Target Plan and Assumptions. Estimated columns are marked "(EST)".
"""
from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..constants import (APP_NAME, APP_VERSION, ESTIMATE_DISCLAIMER, FEE_DISCLAIMER, HIGH_RISK, INCOMPLETE,
                         NEEDS_VERIFICATION, QUALIFIED, REJECTED, RESTRICTION_LABELS, STATUS_LABELS)
from ..domain.settings import FIELD_SPECS, SETTING_GROUPS, VAT_MODE_LABELS, Settings
from ..fees.engine import load_fee_table
from ..formatting import parse_iso
from .ranking import rank

INK = "23262E"
GBP = '£#,##0.00;[Red]-£#,##0.00'
PCT = '0.0%'
INT = '#,##0'
DATE = 'dd mmm yyyy'
TONES = {
    "ok": ("E7F6EE", "047857"),
    "warn": ("FFF3DF", "B45309"),
    "bad": ("FDECEC", "B91C1C"),
    "mute": ("EEF1F6", "4B5563"),
}
RISK_TONE = {"LOW": "ok", "MEDIUM": "warn", "HIGH": "bad"}
STATUS_TONE = {QUALIFIED: "ok", NEEDS_VERIFICATION: "warn", HIGH_RISK: "bad", REJECTED: "bad", INCOMPLETE: "mute"}
THIN = Side(style="thin", color="E3E7EE")


def safe_text(value):
    """Stop spreadsheet formula injection from user-entered text."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _date(value):
    dt = parse_iso(value)
    return dt.replace(tzinfo=None) if dt else None


def _pct(value):
    return None if value is None else value / 100


def _surcharges(fees):
    if fees["total_fees"] is None:
        return None
    return round(fees["fuel_surcharge"] + fees["dangerous_goods_fee"] + fees["digital_services_fee"] + fees["vat_on_fees"], 2)


def _risk_reasons(ev):
    return "; ".join(f"{r['label']} (+{r['points']})" for r in ev["risk"]["reasons"]) or "None"


# (header, width, number format, value getter(item, ctx))
def _core_columns():
    return [
        ("ASIN", 13, None, lambda i, c: i["product"].asin),
        ("Product", 44, None, lambda i, c: i["product"].display_name),
        ("Brand", 18, None, lambda i, c: i["product"].brand),
        ("Category", 20, None, lambda i, c: i["product"].category),
        ("Amazon price", 12, GBP, lambda i, c: i["product"].selling_price),
        ("BSR", 10, INT, lambda i, c: i["product"].bsr),
        ("Sellers", 9, INT, lambda i, c: i["product"].total_sellers),
        ("FBA sellers", 9, INT, lambda i, c: i["product"].fba_sellers),
        ("Estimated monthly sales (EST)", 13, INT, lambda i, c: i["product"].monthly_sales),
        ("Sourcing price", 12, GBP, lambda i, c: i["product"].sourcing_price),
        ("Supplier", 20, None, lambda i, c: i["product"].supplier),
        ("Referral fee (EST)", 12, GBP, lambda i, c: i["evaluation"]["fees"]["referral_fee"]),
        ("FBA fee (EST)", 11, GBP, lambda i, c: i["evaluation"]["fees"]["fba_fee"]),
        ("Surcharges, DSF & VAT on fees (EST)", 14, GBP, lambda i, c: _surcharges(i["evaluation"]["fees"])),
        ("Total Amazon fees (EST)", 12, GBP, lambda i, c: i["evaluation"]["fees"]["total_fees"]),
        ("Shipping/prep", 11, GBP, lambda i, c: i["product"].shipping_prep),
        ("Other costs", 11, GBP, lambda i, c: i["product"].other_costs),
        ("Output VAT (EST)", 11, GBP, lambda i, c: i["evaluation"]["profit"]["output_vat"] or None),
        ("Net profit/unit (EST)", 12, GBP, lambda i, c: i["evaluation"]["profit"]["net_profit"]),
        ("ROI (EST)", 9, PCT, lambda i, c: _pct(i["evaluation"]["profit"]["roi_pct"])),
        ("Capture rate (EST)", 10, PCT, lambda i, c: _pct(i["evaluation"]["profit"]["capture_rate_pct"])),
        ("Conservative units/month (EST)", 13, INT, lambda i, c: i["evaluation"]["profit"]["conservative_units"]),
        ("Conservative monthly profit (EST)", 14, GBP, lambda i, c: i["evaluation"]["profit"]["conservative_monthly_profit"]),
        ("Theoretical monthly profit (EST)", 14, GBP, lambda i, c: i["evaluation"]["profit"]["theoretical_monthly_profit"]),
        ("Restriction status", 16, None, lambda i, c: RESTRICTION_LABELS.get(i["product"].restriction_status)),
        ("Risk", 9, None, lambda i, c: i["evaluation"]["risk"]["level"]),
        ("Risk reasons", 44, None, lambda i, c: _risk_reasons(i["evaluation"])),
    ]


def _tail_columns():
    return [
        ("Source URL", 34, None, lambda i, c: i["product"].amazon_url),
        ("Checked date", 13, DATE, lambda i, c: _date(i["product"].amazon_checked_at)),
    ]


def _write_table(ws, columns, items, ctx, empty_text, freeze="C2"):
    ws.append([c[0] for c in columns])
    header_fill = PatternFill("solid", fgColor=INK)
    for idx, (title, width, _, _) in enumerate(columns, start=1):
        cell = ws.cell(1, idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 42

    if not items:
        ws.append([empty_text])
        ws.cell(2, 1).font = Font(italic=True, color="6B7280")
    for item in items:
        ws.append([safe_text(getter(item, ctx)) for _, _, _, getter in columns])
        row = ws.max_row
        for idx, (title, _, fmt, _) in enumerate(columns, start=1):
            cell = ws.cell(row, idx)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=title in ("Product", "Risk reasons", "Reasons",
                                                                             "Rejected because", "Missing data"))
            if fmt:
                cell.number_format = fmt
            if title == "Risk" and cell.value in RISK_TONE:
                _tone(cell, RISK_TONE[cell.value])
            if title == "Status":
                status = next((k for k, v in STATUS_LABELS.items() if v == cell.value), None)
                if status:
                    _tone(cell, STATUS_TONE[status])
            if title == "Included in target plan" and cell.value == "Yes":
                _tone(cell, "ok")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(ws.max_row, 2)}"
    ws.sheet_view.zoomScale = 90


def _tone(cell, tone):
    fill, color = TONES[tone]
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(bold=True, color=color)


def build_workbook(items: list[dict], plan: dict, settings: Settings, generated_at: datetime | None = None) -> bytes:
    generated_at = generated_at or datetime.now()
    ctx = {"plan_ids": plan["selected_ids"]}
    by_status = {s: [i for i in items if i["evaluation"]["status"] == s] for s in STATUS_LABELS}
    qualified = rank(items)

    wb = Workbook()
    ws = wb.active
    ws.title = "Qualified"
    q_cols = ([("Rank", 7, INT, lambda i, c: i.get("rank")), ("Score", 8, INT,
               lambda i, c: (i["evaluation"].get("score") or {}).get("total"))]
              + _core_columns()
              + [("Included in target plan", 12, None, lambda i, c: "Yes" if i["product"].id in c["plan_ids"] else "No")]
              + _tail_columns())
    _write_table(ws, q_cols, qualified, ctx, "No qualified products yet.", freeze="E2")

    ws = wb.create_sheet("Rejected")
    r_cols = _core_columns() + [("Rejected because", 50, None, lambda i, c: "\n".join(i["evaluation"]["summary"]))] + _tail_columns()
    _write_table(ws, r_cols, sorted(by_status[REJECTED], key=lambda i: i["product"].display_name.lower()), ctx,
                 "No rejected products.")

    ws = wb.create_sheet("Incomplete")
    i_cols = _core_columns()[:11] + [
        ("Missing data", 60, None, lambda i, c: "\n".join(m["message"] for m in i["evaluation"]["missing"])),
    ] + _tail_columns()
    _write_table(ws, i_cols, sorted(by_status[INCOMPLETE], key=lambda i: i["product"].display_name.lower()), ctx,
                 "No products with missing data.")

    _summary_sheet(wb.create_sheet("Summary"), items, by_status, plan, settings, generated_at)

    ws = wb.create_sheet("Needs Review")
    n_cols = [("Status", 18, None, lambda i, c: i["evaluation"]["status_label"])] + _core_columns() + [
        ("Reasons", 50, None, lambda i, c: "\n".join(i["evaluation"]["summary"]))] + _tail_columns()
    review = sorted(by_status[HIGH_RISK] + by_status[NEEDS_VERIFICATION], key=lambda i: i["product"].display_name.lower())
    _write_table(ws, n_cols, review, ctx, "Nothing needs review.", freeze="D2")

    _plan_sheet(wb.create_sheet("Target Plan"), plan)
    _assumptions_sheet(wb.create_sheet("Assumptions"), settings)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _summary_sheet(ws, items, by_status, plan, settings, generated_at):
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22
    ws.append([f"{APP_NAME} — Summary"])
    ws["A1"].font = Font(bold=True, size=15)
    ws.append(["ESTIMATES ONLY. " + ESTIMATE_DISCLAIMER + " " + FEE_DISCLAIMER])
    ws["A2"].font = Font(italic=True, color="B45309")
    ws.append([])
    qualified = by_status[QUALIFIED]
    cons = round(sum(i["evaluation"]["profit"]["conservative_monthly_profit"] or 0 for i in qualified), 2)
    theo = round(sum(i["evaluation"]["profit"]["theoretical_monthly_profit"] or 0 for i in qualified), 2)
    rows = [
        ("Product count", len(items), INT),
        ("Qualified count", len(qualified), INT),
        ("Rejected count", len(by_status[REJECTED]), INT),
        ("Incomplete count", len(by_status[INCOMPLETE]), INT),
        ("Needs Verification count", len(by_status[NEEDS_VERIFICATION]), INT),
        ("High Risk count", len(by_status[HIGH_RISK]), INT),
        ("Conservative monthly profit — all qualified (EST)", cons, GBP),
        ("Theoretical monthly profit — all qualified (EST)", theo, GBP),
        ("Monthly profit target", plan["target"], GBP),
        ("Target plan total — conservative (EST)", plan["total"], GBP),
        ("Target achieved?", "YES" if plan["achieved"] else "NO", None),
        ("Products in target plan", len(plan["selected"]), INT),
        ("Monthly stock investment for the plan (EST)", plan["stock_investment"], GBP),
        ("Date generated", generated_at.replace(microsecond=0), "dd mmm yyyy hh:mm"),
        ("Fee table version", load_fee_table()["version"], None),
        ("App version", APP_VERSION, None),
    ]
    for label, value, fmt in rows:
        ws.append([label, value])
        cell = ws.cell(ws.max_row, 2)
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right")
        if fmt:
            cell.number_format = fmt
        if label == "Target achieved?":
            _tone(cell, "ok" if plan["achieved"] else "warn")


def _plan_sheet(ws, plan):
    cols = [("Rank", 7), ("Product", 46), ("ASIN", 13), ("Score", 8), ("Conservative units/month (EST)", 14),
            ("Conservative monthly profit (EST)", 16), ("Running total (EST)", 16), ("Monthly stock investment (EST)", 16)]
    ws.append([c[0] for c in cols])
    for idx, (title, width) in enumerate(cols, start=1):
        cell = ws.cell(1, idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 42
    for s in plan["selected"]:
        ws.append([s["rank"], safe_text(s["name"]), s["asin"], s["score"], s["conservative_units"],
                   s["conservative_monthly_profit"], s["running_total"], s["monthly_stock_investment"]])
        for col, fmt in ((5, INT), (6, GBP), (7, GBP), (8, GBP)):
            ws.cell(ws.max_row, col).number_format = fmt
    if not plan["selected"]:
        ws.append(["No qualified products with positive conservative profit yet."])
    ws.append([])
    for label, value, fmt in (("Target", plan["target"], GBP), ("Projected conservative total (EST)", plan["total"], GBP),
                              ("Target achieved?", "YES" if plan["achieved"] else "NO", None),
                              ("Shortfall", plan["shortfall"], GBP)):
        ws.append(["", label, "", "", "", value])
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        if fmt:
            ws.cell(ws.max_row, 6).number_format = fmt
    ws.append([])
    ws.append(["", "Estimated, not guaranteed income. Conservative figures assume you win only the capture rate "
                   "share of the listing's estimated sales."])
    ws.cell(ws.max_row, 2).font = Font(italic=True, color="B45309")
    ws.freeze_panes = "A2"


def _assumptions_sheet(ws, settings: Settings):
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 34
    ws.append(["Assumptions used for this export"])
    ws["A1"].font = Font(bold=True, size=14)
    for group in SETTING_GROUPS:
        ws.append([])
        ws.append([group["title"] + (" (ESTIMATED)" if group.get("estimated") else "")])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for spec in group["fields"]:
            value = getattr(settings, spec["name"])
            if spec["kind"] == "bool":
                value = "Yes" if value else "No"
            elif spec["kind"] == "choice":
                value = VAT_MODE_LABELS.get(value, value)
            elif spec["kind"] == "pct":
                value = f"{value:g}%"
            elif spec["kind"] == "money":
                value = f"£{value:,.2f}"
            ws.append([FIELD_SPECS[spec["name"]]["label"], value])
    table = load_fee_table()
    ws.append([])
    ws.append(["Fee table"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["Version", table["version"]])
    ws.append(["Effective from", table["effective_from"]])
    for source in table["sources"]:
        ws.append(["Source", source])
    ws.append([FEE_DISCLAIMER])
    ws.cell(ws.max_row, 1).font = Font(italic=True, color="B45309")


CSV_COLUMNS = ["Status", "Score", "ASIN", "Product", "Brand", "Category", "Amazon price", "BSR", "Sellers", "FBA sellers",
               "Estimated monthly sales (EST)", "Sourcing price", "Supplier", "Total Amazon fees (EST)", "Net profit/unit (EST)",
               "ROI % (EST)", "Conservative units/month (EST)", "Conservative monthly profit (EST)",
               "Theoretical monthly profit (EST)", "Risk", "Restriction status", "Reasons", "Amazon URL", "Last checked"]


def build_csv(items: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(CSV_COLUMNS)
    for i in items:
        p, e = i["product"], i["evaluation"]
        pr = e["profit"]
        writer.writerow([safe_text(v) if isinstance(v, str) else ("" if v is None else v) for v in [
            e["status_label"], (e.get("score") or {}).get("total"), p.asin, p.display_name, p.brand, p.category,
            p.selling_price, p.bsr, p.total_sellers, p.fba_sellers, p.monthly_sales, p.sourcing_price, p.supplier,
            e["fees"]["total_fees"], pr["net_profit"], pr["roi_pct"], pr["conservative_units"],
            pr["conservative_monthly_profit"], pr["theoretical_monthly_profit"], e["risk"]["level"],
            RESTRICTION_LABELS.get(p.restriction_status), " | ".join(e["summary"]), p.amazon_url, p.amazon_checked_at,
        ]])
    return ("﻿" + buf.getvalue()).encode("utf-8")
