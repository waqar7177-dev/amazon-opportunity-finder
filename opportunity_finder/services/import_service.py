"""Bulk import from CSV, XLSX or pasted tab-separated text.

Flow: read -> auto-map columns -> preview (every row validated, duplicates found) -> commit.
Nothing is written until the user confirms, and the commit runs in one transaction,
so a bad file can never leave half an import behind.
"""
from __future__ import annotations

import csv
import io
import json
import re
import time
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from ..constants import FLAG_KEYS, SOURCE_IMPORT
from ..domain.product import Product
from ..numbers import is_blank
from .duplicates import DuplicateMatch, find_duplicates, identity_key
from .validation import validate_product

MAX_ROWS = 5000
SESSION_TTL_SECONDS = 24 * 3600

# field key -> (column title used in the template, accepted header spellings)
IMPORT_FIELDS: list[tuple[str, str, list[str]]] = [
    ("asin", "ASIN", ["asin"]),
    ("title", "Title", ["title", "product", "product title", "product name", "name", "item name"]),
    ("brand", "Brand", ["brand", "brand name"]),
    ("category", "Category", ["category", "amazon category", "bsr category"]),
    ("fee_category", "Fee category", ["fee category", "referral category", "referral fee category"]),
    ("amazon_url", "Amazon URL", ["amazon url", "url", "link", "product url", "source url", "amazon link"]),
    ("selling_price", "Amazon price", ["amazon price", "selling price", "price", "buy box", "buy box price", "sell price",
                                       "sale price", "amazon selling price"]),
    ("bsr", "BSR", ["bsr", "sales rank", "best sellers rank", "best seller rank", "rank"]),
    ("monthly_sales", "Estimated monthly sales", ["estimated monthly sales", "monthly sales", "est monthly sales",
                                                  "sales per month", "units per month", "monthly units", "sales"]),
    ("total_sellers", "Total sellers", ["total sellers", "sellers", "seller count", "offers", "offer count", "number of sellers"]),
    ("fba_sellers", "FBA sellers", ["fba sellers", "fba offers", "fba seller count", "fba offer count"]),
    ("sourcing_price", "Sourcing price", ["sourcing price", "cost", "buy cost", "unit cost", "cost price", "supplier price",
                                          "wholesale price", "buy price", "cost per unit"]),
    ("supplier", "Supplier", ["supplier", "vendor", "source", "supplier name"]),
    ("shipping_prep", "Shipping/prep", ["shipping prep", "shipping", "prep", "prep cost", "inbound shipping",
                                        "shipping cost", "shipping and prep"]),
    ("other_costs", "Other costs", ["other costs", "other cost", "other", "misc", "misc costs"]),
    ("weight_g", "Weight (g)", ["weight g", "weight", "weight grams", "weight in grams", "package weight g"]),
    ("weight_kg", "Weight (kg)", ["weight kg", "weight in kg", "package weight kg"]),
    ("length_cm", "Length (cm)", ["length cm", "length"]),
    ("width_cm", "Width (cm)", ["width cm", "width"]),
    ("height_cm", "Height (cm)", ["height cm", "height", "depth cm", "depth"]),
    ("fulfilment", "Fulfilment", ["fulfilment", "fulfillment", "fulfilment method", "fulfillment method", "channel"]),
    ("restriction_status", "Restriction status", ["restriction status", "restriction", "restrictions", "gated", "gating"]),
    ("flag_hazmat", "Hazmat", ["hazmat", "dangerous goods", "flag hazmat"]),
    ("flag_battery", "Battery", ["battery", "batteries", "flag battery"]),
    ("flag_liquid", "Liquid", ["liquid", "flag liquid"]),
    ("flag_fragile", "Fragile", ["fragile", "flag fragile"]),
    ("flag_seasonal", "Seasonal", ["seasonal", "flag seasonal"]),
    ("flag_amazon_sells", "Amazon sells", ["amazon sells", "amazon on listing", "amazon seller", "amazon selling",
                                           "flag amazon sells"]),
    ("flag_ip_concern", "IP concern", ["ip concern", "ip risk", "trademark concern", "flag ip concern"]),
    ("referral_fee_override", "Referral fee override", ["referral fee override", "referral fee"]),
    ("fba_fee_override", "FBA fee override", ["fba fee override", "fba fee", "fulfilment fee", "fulfillment fee"]),
    ("notes", "Notes", ["notes", "note", "comments", "comment"]),
]
FIELD_TITLES = {key: title for key, title, _ in IMPORT_FIELDS}
IDENTITY_FIELDS = ("asin", "title", "amazon_url")
DUPLICATE_POLICIES = {
    "skip": "Skip rows that match a product you already have",
    "update": "Update the existing product with the row's values",
    "create": "Create a new product anyway",
}


class ImportFileError(ValueError):
    """A problem with the file as a whole — shown to the user as-is."""


@dataclass
class RawTable:
    headers: list[str]
    rows: list[list]
    source_name: str
    blank_rows: int = 0

    def to_json(self) -> dict:
        return {"headers": self.headers, "rows": self.rows, "source_name": self.source_name, "blank_rows": self.blank_rows}

    @classmethod
    def from_json(cls, data: dict) -> "RawTable":
        return cls(data["headers"], data["rows"], data["source_name"], data.get("blank_rows", 0))


@dataclass
class RowResult:
    row_number: int
    raw: dict
    product: Product | None
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    duplicate: DuplicateMatch | None = None
    in_file_duplicate_of: int | None = None
    action: str = "import"  # import | update | skip | invalid
    provided: set = field(default_factory=set)


@dataclass
class Preview:
    rows: list[RowResult]
    mapping: dict[int, str]
    mapping_errors: list[str]
    policy: str

    def count(self, action: str) -> int:
        return sum(1 for r in self.rows if r.action == action)

    @property
    def counts(self) -> dict:
        return {a: self.count(a) for a in ("import", "update", "skip", "invalid")} | {"total": len(self.rows)}


@dataclass
class ImportReport:
    imported: int = 0
    updated: int = 0
    skipped: int = 0
    failed: list[tuple[int, list[str]]] = field(default_factory=list)
    new_ids: list[int] = field(default_factory=list)


# ---------------------------------------------------------------- reading
def _normalise_header(text) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower().replace("£", "")).strip()


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, (int, float)):
        return value
    return str(value).strip()


def _finish(rows: list[list], source_name: str) -> RawTable:
    rows = [[_cell(c) for c in row] for row in rows]
    while rows and all(c == "" for c in rows[0]):
        rows.pop(0)
    if not rows:
        raise ImportFileError("The file is empty.")
    headers = [str(h).strip() for h in rows[0]]
    while headers and headers[-1] == "":
        headers.pop()
    if not any(headers):
        raise ImportFileError("The first row must contain column headings.")
    data, blank = [], 0
    for row in rows[1:]:
        row = (row + [""] * len(headers))[: len(headers)]
        if all(c == "" for c in row):
            blank += 1
            continue
        data.append(row)
    if not data:
        raise ImportFileError("No data rows were found under the heading row.")
    if len(data) > MAX_ROWS:
        raise ImportFileError(f"The file has {len(data):,} rows. Import at most {MAX_ROWS:,} rows at a time.")
    return RawTable(headers, data, source_name, blank)


def read_csv_bytes(data: bytes, source_name: str = "upload.csv") -> RawTable:
    if not data or not data.strip():
        raise ImportFileError("The file is empty.")
    if b"\x00" in data[:4096]:
        raise ImportFileError("This doesn't look like a CSV text file. Save the sheet as CSV (UTF-8) or upload the .xlsx.")
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    return read_delimited_text(text, source_name)


def read_delimited_text(text: str, source_name: str = "pasted data") -> RawTable:
    if not text or not text.strip():
        raise ImportFileError("Nothing was pasted.")
    first_line = text.strip().splitlines()[0]
    if "\t" in first_line:
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
    try:
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error as exc:
        raise ImportFileError(f"The data could not be read as a table ({exc}).") from exc
    return _finish(rows, source_name)


def read_xlsx_bytes(data: bytes, source_name: str = "upload.xlsx") -> RawTable:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError) as exc:
        raise ImportFileError("This isn't a valid .xlsx file. Open it in Excel and save it as .xlsx or CSV.") from exc
    try:
        sheet = wb["Products"] if "Products" in wb.sheetnames else wb.worksheets[0]
        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i > MAX_ROWS + 50:
                raise ImportFileError(f"The sheet has too many rows. Import at most {MAX_ROWS:,} rows at a time.")
            rows.append(list(row))
    finally:
        wb.close()
    return _finish(rows, source_name)


def read_upload(filename: str, data: bytes) -> RawTable:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return read_xlsx_bytes(data, filename)
    if name.endswith(".xls"):
        raise ImportFileError("Old .xls files aren't supported. In Excel choose File → Save As → .xlsx or CSV.")
    if name.endswith((".csv", ".txt", ".tsv")):
        return read_csv_bytes(data, filename)
    if data[:2] == b"PK":
        return read_xlsx_bytes(data, filename)
    raise ImportFileError("Upload a .csv or .xlsx file.")


# ---------------------------------------------------------------- mapping
def auto_map(headers: list[str]) -> dict[int, str]:
    lookup = {}
    for key, title, aliases in IMPORT_FIELDS:
        for alias in [title, key.replace("_", " "), *aliases]:
            lookup.setdefault(_normalise_header(alias), key)
    mapping, used = {}, set()
    for index, header in enumerate(headers):
        key = lookup.get(_normalise_header(header))
        if key and key not in used:
            mapping[index] = key
            used.add(key)
    return mapping


def mapping_from_form(form, headers: list[str]) -> dict[int, str]:
    valid = set(FIELD_TITLES)
    mapping = {}
    for index in range(len(headers)):
        value = form.get(f"map_{index}", "")
        if value in valid:
            mapping[index] = value
    return mapping


def check_mapping(mapping: dict[int, str]) -> list[str]:
    errors = []
    values = list(mapping.values())
    dupes = sorted({v for v in values if values.count(v) > 1})
    for key in dupes:
        errors.append(f"“{FIELD_TITLES[key]}” is mapped to more than one column. Map each field once.")
    if not set(values) & set(IDENTITY_FIELDS):
        errors.append("Map at least one of ASIN, Title or Amazon URL so each product can be recognised.")
    if "weight_g" in values and "weight_kg" in values:
        errors.append("Map either Weight (g) or Weight (kg), not both.")
    return errors


# ---------------------------------------------------------------- preview
def build_preview(table: RawTable, mapping: dict[int, str], repo, policy: str = "skip") -> Preview:
    policy = policy if policy in DUPLICATE_POLICIES else "skip"
    mapping_errors = check_mapping(mapping)
    rows: list[RowResult] = []
    seen: dict[str, int] = {}
    for offset, raw_row in enumerate(table.rows):
        row_number = offset + 2  # spreadsheet row number (row 1 is the header)
        data = {key: raw_row[index] for index, key in mapping.items() if index < len(raw_row)}
        provided = {k for k, v in data.items() if not is_blank(v)} | {k for k in data if k in FLAG_KEYS}
        if "weight_kg" in provided:
            provided.add("weight_g")
        result = RowResult(row_number, {FIELD_TITLES[k]: v for k, v in data.items()}, None, provided=provided)
        if mapping_errors:
            result.action = "invalid"
            rows.append(result)
            continue
        data["source"] = SOURCE_IMPORT
        validation = validate_product(data, checkbox_form=False)
        result.product = validation.product
        result.errors = list(validation.errors.values())
        result.warnings = list(validation.warnings.values())
        if result.errors:
            result.action = "invalid"
            rows.append(result)
            continue
        key = identity_key(validation.product)
        if key in seen:
            result.in_file_duplicate_of = seen[key]
            result.errors.append(f"Same product as row {seen[key]} in this file.")
            result.action = "invalid"
            rows.append(result)
            continue
        if key:
            seen[key] = row_number
        matches = find_duplicates(repo, validation.product)
        if matches:
            result.duplicate = matches[0]
            result.action = {"skip": "skip", "update": "update", "create": "import"}[policy]
        rows.append(result)
    return Preview(rows, mapping, mapping_errors, policy)


def commit_preview(preview: Preview, service) -> ImportReport:
    """Write the valid rows in one transaction. Any database error rolls everything back."""
    report = ImportReport()
    if preview.mapping_errors:
        raise ImportFileError(preview.mapping_errors[0])
    conn = service.conn
    try:
        for row in preview.rows:
            if row.action == "invalid":
                report.failed.append((row.row_number, row.errors or ["Invalid row."]))
            elif row.action == "skip":
                report.skipped += 1
            elif row.action == "update":
                service.merge(row.duplicate.product.id, row.product, provided=row.provided, source=SOURCE_IMPORT, commit=False)
                report.updated += 1
            else:
                report.new_ids.append(service.create(row.product, source=SOURCE_IMPORT, commit=False))
                report.imported += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    return report


# ---------------------------------------------------------------- sessions
def save_session(import_dir: Path, table: RawTable) -> str:
    import_dir.mkdir(parents=True, exist_ok=True)
    cleanup_sessions(import_dir)
    token = uuid.uuid4().hex
    (import_dir / f"{token}.json").write_text(json.dumps(table.to_json()), encoding="utf-8")
    return token


def load_session(import_dir: Path, token: str) -> RawTable | None:
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        return None
    path = import_dir / f"{token}.json"
    if not path.exists():
        return None
    try:
        return RawTable.from_json(json.loads(path.read_text(encoding="utf-8")))
    except (ValueError, KeyError):
        return None


def delete_session(import_dir: Path, token: str) -> None:
    if re.fullmatch(r"[0-9a-f]{32}", token or ""):
        (import_dir / f"{token}.json").unlink(missing_ok=True)


def cleanup_sessions(import_dir: Path) -> None:
    cutoff = time.time() - SESSION_TTL_SECONDS
    for path in import_dir.glob("*.json"):
        try:
            if path.stat().st_mtime < cutoff:
                path.unlink()
        except OSError:
            pass


# ---------------------------------------------------------------- templates
TEMPLATE_FIELDS = [k for k, _, _ in IMPORT_FIELDS if k != "weight_kg"]
EXAMPLE_ROW = {
    "asin": "B0EXAMPLE1", "title": "EXAMPLE ROW - replace with your product", "brand": "Example Brand",
    "category": "Home & Kitchen", "fee_category": "Everything else", "amazon_url": "https://www.amazon.co.uk/dp/B0EXAMPLE1",
    "selling_price": 19.99, "bsr": 8500, "monthly_sales": 240, "total_sellers": 5, "fba_sellers": 3,
    "sourcing_price": 6.50, "supplier": "Example Wholesale Ltd", "shipping_prep": 0.45, "other_costs": 0,
    "weight_g": 350, "length_cm": 22, "width_cm": 16, "height_cm": 6, "fulfilment": "FBA",
    "restriction_status": "Unknown", "flag_hazmat": "No", "flag_battery": "No", "flag_liquid": "No", "flag_fragile": "No",
    "flag_seasonal": "No", "flag_amazon_sells": "No", "flag_ip_concern": "No", "referral_fee_override": "",
    "fba_fee_override": "", "notes": "Example only - delete this row",
}


def template_csv() -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([FIELD_TITLES[k] for k in TEMPLATE_FIELDS])
    writer.writerow([EXAMPLE_ROW[k] for k in TEMPLATE_FIELDS])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from ..fees.engine import referral_categories

    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append([FIELD_TITLES[k] for k in TEMPLATE_FIELDS])
    ws.append([EXAMPLE_ROW[k] for k in TEMPLATE_FIELDS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="23262E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for idx, key in enumerate(TEMPLATE_FIELDS, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = max(12, min(40, len(FIELD_TITLES[key]) + 6))
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instructions")
    lines = [
        ("How to fill in this template", True),
        ("One product per row on the Products sheet. Delete the example row first.", False),
        ("Required: at least one of ASIN, Title or Amazon URL. Everything else is optional —", False),
        ("missing values are never guessed; the product is marked Incomplete Data until you add them.", False),
        ("Prices and costs in GBP, e.g. 12.99 (a £ sign and commas are fine).", False),
        ("BSR, monthly sales and seller counts must be whole numbers.", False),
        ("Fulfilment: FBA or FBM. Restriction status: Unknown, Check Seller Central, Restricted or Ungated.", False),
        ("Yes/No columns accept Yes, No, Y, N, 1, 0, True, False.", False),
        ("Fee category: use a name from the Fee categories sheet (blank = Everything else, 15%).", False),
        ("Referral / FBA fee override: only if you have the exact fee from Seller Central.", False),
    ]
    for text, bold in lines:
        info.append([text])
        if bold:
            info.cell(info.max_row, 1).font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 110

    cats = wb.create_sheet("Fee categories")
    cats.append(["Fee category", "Key"])
    for c in cats[1]:
        c.font = Font(bold=True)
    for cat in referral_categories():
        cats.append([cat["label"], cat["key"]])
    cats.column_dimensions["A"].width = 48
    cats.column_dimensions["B"].width = 34

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
